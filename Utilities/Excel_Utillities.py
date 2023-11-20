import os
import openpyxl

relative = "./TestData/Data_Driven_Test.xlsx"
path = os.path.abspath(relative)
workbook = openpyxl.load_workbook(path)


def get_row_count(sheet_name):
    sheet = workbook[sheet_name]
    return sheet.max_row


def get_column_count(sheet_name):
    sheet = workbook[sheet_name]
    return sheet.max_column


def read_data(sheet_name, row_num, column_num):
    sheet = workbook[sheet_name]
    return sheet.cell(row=row_num, column=column_num).value


def write_data(sheet_name, row_num, column_num, data):
    sheet = workbook[sheet_name]
    sheet.cell(row=row_num, column=column_num).value = data
    workbook.save(path)
    